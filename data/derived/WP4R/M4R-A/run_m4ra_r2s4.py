#!/usr/bin/env python
"""M4R-A ROUND2 step4: perturbation crosswalk (PfAP2-P DE + bound x Asn axis; PbApiAP2 inventory)."""
import glob, os, re, zipfile
import numpy as np
import pandas as pd
import openpyxl
from scipy import stats

ROOT = "/home/huyudi/015_plasmo"
OUT = f"{ROOT}/data/derived/WP4R/M4R-A"
AP2 = f"{ROOT}/data/manual_inbox/pmc10627835_tmp"
PB = f"{ROOT}/data/manual_inbox/M4R_PBAPIAP2_2017/content"

def gkey(s):
    m = re.search(r"PF3D7_\d+", str(s)); return m.group(0) if m else None

def mwu(a, b):
    a = np.asarray(a, float); b = np.asarray(b, float)
    a = a[np.isfinite(a)]; b = b[np.isfinite(b)]
    if len(a) < 3 or len(b) < 3: return float("nan"), float("nan"), float("nan"), float("nan")
    u, p = stats.mannwhitneyu(a, b, alternative="two-sided", method="asymptotic")
    return float(u), float(p), float(np.median(a)), float(np.median(b))

# universe features
FAA = f"{ROOT}/data/raw/veupathdb/PlasmoDB-71/PlasmoDB-71_Pfalciparum3D7_AnnotatedProteins.fasta"
prots = {}
with open(FAA) as fh:
    hid, seq = None, []
    def flush():
        if hid is None: return
        g = gkey(hid)
        if not g: return
        s = "".join(seq).upper().replace("*", "")
        if len(s) > len(prots.get(g, "")): prots[g] = s
    for line in fh:
        if line.startswith(">"): flush(); hid, seq = line.rstrip(), []
        else: seq.append(line.strip())
    flush()
FE = pd.DataFrame({"gene": list(prots),
                   "asn_frac": [s.count("N") / len(s) for s in prots.values()]}).set_index("gene")
FE["asn_top10"] = FE.asn_frac >= FE.asn_frac.quantile(0.9)
M = pd.read_csv(f"{OUT}/M4RA_interpro_map.tsv", sep="\t")
g2ipr = dict(zip(M.gene, M.interpro.fillna("").apply(lambda s: set(s.split(";")) if s else set())))
FE["AP2_ip"] = FE.index.map(lambda g: "IPR001471" in g2ipr.get(g, set()))
print("universe:", len(FE), "AP2_ip:", int(FE.AP2_ip.sum()), flush=True)

XW = []
# ---- PfAP2-P MOESM3: DE per stage ----
wb = openpyxl.load_workbook(f"{AP2}/41564_2023_1497_MOESM3_ESM.xlsx", read_only=True, data_only=True)
for sn in ["16 h.p.i.", "40 h.p.i.", "8 h.p.i.", "30 h.p.i."]:
    if sn not in wb.sheetnames: continue
    ws = wb[sn]; rows = list(ws.values)
    hdr = [str(v) for v in rows[0]]
    gi = next(i for i, h in enumerate(hdr) if "Gene" in h)
    lfc = next(i for i, h in enumerate(hdr) if "log2FoldChange" in h)
    pa = next(i for i, h in enumerate(hdr) if h == "padj")
    rec = [(gkey(r[gi]), r[lfc], r[pa]) for r in rows[1:] if gkey(r[gi])]
    d = pd.DataFrame(rec, columns=["gene", "lfc", "padj"]).dropna()
    d["lfc"] = pd.to_numeric(d["lfc"], errors="coerce")
    d["padj"] = pd.to_numeric(d["padj"], errors="coerce")
    d = d.join(FE, on="gene", how="inner")
    print(sn, "de-rows:", len(d), flush=True)
    r, p = stats.spearmanr(d.asn_frac, d.lfc)
    XW.append({"source": "PfAP2-P_trunc_DE", "contrast": sn, "n": len(d),
               "test": "spearman(asn,DElogFC)", "stat": float(r), "p": float(p),
               "note": ""})
    ds = d[d.padj < 0.05]
    u, p2, mh, mr = mwu(ds[ds.lfc > 0].asn_frac, ds[ds.lfc < 0].asn_frac) if len(ds) else (float("nan"),)*4
    XW.append({"source": "PfAP2-P_trunc_DE", "contrast": sn, "n": len(ds),
               "test": "MW Asn(UPpadj.05 vs DOWNpadj.05)", "stat": u, "p": p2,
               "note": f"up_med={mh:.3f} down_med={mr:.3f}" if mh == mh else ""})

# ---- PfAP2-P MOESM5: bound genes ----
wb5 = openpyxl.load_workbook(f"{AP2}/41564_2023_1497_MOESM5_ESM.xlsx", read_only=True, data_only=True)
for sn in ["16 h.p.i.", "40 h.p.i."]:
    ws = wb5[sn]; rows = list(ws.values)
    hdr = [str(v) for v in rows[0]]
    gi = next(i for i, h in enumerate(hdr) if "Gene" in h)
    bg = {gkey(r[gi]) for r in rows[1:] if gkey(r[gi])}
    d = FE.copy(); d["bound"] = d.index.isin(bg)
    u, p, mh, mr = mwu(d[d.bound].asn_frac, d[~d.bound].asn_frac)
    XW.append({"source": "PfAP2-P_bound", "contrast": sn, "n": len(d),
               "test": "MW Asn(bound vs unbound)", "stat": u, "p": p,
               "note": f"n_bound={len(bg)} med {mh:.4f} vs {mr:.4f}"})
    for col, nm in [("asn_top10", "bound x Asn-top10"), ("AP2_ip", "bound x AP2(IPR001471)")]:
        kk = d[col]
        a = int((d.bound & kk).sum()); b = int(((~d.bound) & kk).sum())
        c = int((d.bound & ~kk).sum()); dd = int(((~d.bound) & ~kk).sum())
        o, pp = stats.fisher_exact([[a, b], [c, dd]])
        XW.append({"source": "PfAP2-P_bound", "contrast": sn, "n": len(d),
                   "test": f"Fisher {nm}", "stat": float(o), "p": float(pp),
                   "note": f"overlap={a}"})

X = pd.DataFrame(XW)
X.to_csv(f"{OUT}/M4RA_perturbation_xwalk.tsv", sep="\t", index=False)
print(X.to_string(), flush=True)

# ---- PbApiAP2 inventory (map-only) ----
inv = [{"file": "mmc1.pdf", "kind": "pdf", "content": "unparsed (no pdf tool); main-text supplement assumed"},
       {"file": "mmc7.pdf", "kind": "pdf", "content": "unparsed (no pdf tool); extended supplement assumed"}]
for z, inner in [("mmc2.zip", "Data_1.xlsx"), ("mmc3.zip", "Data_S2.xlsx"),
                 ("mmc4.zip", "Data_S3.xlsx"), ("mmc6.zip", "Data_S5.xlsx")]:
    zp = os.path.join(PB, z)
    with zipfile.ZipFile(zp) as zh:
        names = [n for n in zh.namelist() if not n.startswith("__MACOSX")]
        data = zh.read(names[0])
    with open(f"/tmp/pbinv_{inner}", "wb") as f: f.write(data)
    wb = openpyxl.load_workbook(f"/tmp/pbinv_{inner}", read_only=True, data_only=True)
    for sn in wb.sheetnames[:4]:
        ws = wb[sn]; vals = list(ws.values)
        inv.append({"file": f"{z}:{inner}:{sn}", "kind": "xlsx",
                    "content": f"dims {ws.max_row}x{ws.max_column} | header: {str(vals[0])[:220]}"})
pd.DataFrame(inv).to_csv(f"{OUT}/M4RA_pbapiap2_inventory.tsv", sep="\t", index=False)
print(pd.DataFrame(inv).to_string(), flush=True)
print("STEP4 DONE", flush=True)
