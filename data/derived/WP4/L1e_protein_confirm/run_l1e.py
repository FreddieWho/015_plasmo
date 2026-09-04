#!/usr/bin/env python
# L1e (ring 3): protein-level confirmation of the Asn-rich axis using Small-Saunders 2024 TMT MSstats
# prediction from chain: Asn-rich proteins respond MORE strongly to DHA pulse (up), matching mRNA(L1)+TE(L1-rep)
import hashlib, json, os, re
from collections import Counter
import numpy as np
import pandas as pd
import openpyxl
from scipy import stats

ROOT = "/home/huyudi/015_plasmo"
OUT = f"{ROOT}/data/derived/WP4/L1e_protein_confirm"
os.makedirs(OUT, exist_ok=True)
XLSX = f"{ROOT}/data/manual_inbox/M03_NATMICO2024_SOURCE_DATA/41564_2024_1664_MOESM4_ESM.xlsx"

def sha(p):
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for ch in iter(lambda: f.read(1 << 20), b""): h.update(ch)
    return h.hexdigest()

# ---------- gene features (same CDS rules as L1/L1d) ----------
STD = dict(zip([a + b + c for a in "TCAG" for b in "TCAG" for c in "TCAG"],
               "FFLLSSSSYY**CC*WLLLLPPPPHHQQRRRRIIIMTTTTNNKKSSRRVVVVAAAADDEEGGGG"))
acc = "GCF_000002765.6"
CDS = f"{ROOT}/data/raw/ncbi-datasets/{acc}/ncbi_dataset/data/{acc}/cds_from_genomic.fna"
best = {}
with open(CDS) as fh:
    hid, seq = None, []
    def flush():
        if hid is None: return
        m = re.search(r"\[locus_tag=([^\]]+)\]", hid)
        if not m: return
        g = re.search(r"PF3D7_\d+", m.group(1))
        if not g: return
        g = g.group(0)
        s = "".join(seq).upper().replace("U", "T")
        if "[partial" in hid or len(s) < 150 or len(s) % 3 or not set(s) <= set("ATGC"): return
        cod = [s[i:i+3] for i in range(0, len(s), 3)]
        if "*" in [STD.get(c, "*") for c in cod][:-1]: return
        if len(s) > len(best.get(g, ("", ""))[1]): best[g] = (hid, s)
    for line in fh:
        if line.startswith(">"): flush(); hid, seq = line.rstrip(), []
        else: seq.append(line.strip())
    flush()
feat = {}
for g, (_, s) in best.items():
    cod = [s[i:i+3] for i in range(0, len(s), 3) if STD.get(s[i:i+3], "*") != "*"]
    if not cod: continue
    cc = Counter(cod); n = len(cod)
    asn = cc["AAT"] + cc["AAC"]
    feat[g] = {"AAT_freq": cc["AAT"] / n, "Asn_content": asn / n,
               "AAT_share": cc["AAT"] / asn if asn else np.nan}
FE = pd.DataFrame(feat).T

# products for AP2/PUF flags
prod = {}
for line in open(f"{ROOT}/data/raw/veupathdb/PlasmoDB-71/PlasmoDB-71_Pfalciparum3D7.gff"):
    if "\tprotein_coding_gene\t" not in line: continue
    a = line.split("\t")[8]
    m = re.search(r"ID=(PF3D7_\d+)", a); d = re.search(r"description=([^;]+)", a)
    if m: prod[m.group(1)] = (d.group(1) if d else "").lower()
FE["is_ap2"] = [bool(re.search(r"\bap2|apiap2", prod.get(g, ""))) for g in FE.index]
FE["is_puf"] = [bool(re.search(r"puf", prod.get(g, ""))) for g in FE.index]

# ---------- parse MSstats side-by-side blocks ----------
wb = openpyxl.load_workbook(XLSX, read_only=True)
ws = wb["Source Data Fig. msstats-output"]
rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column, values_only=True))
wb.close()
hdr = rows[1]
# accession -> gene map from first block (cols 0-2)
acc2gene = {}
for r in rows[2:]:
    if r[0] and r[2] and re.match(r"PF3D7", str(r[2])):
        acc2gene[str(r[0])] = re.sub(r"\.\d+$", "", str(r[2]))
# find blocks: col j where hdr[j] in ("Protein Accesion","Protein") and hdr[j+1]=="Label"
blocks = []
for j in range(len(hdr) - 4):
    if str(hdr[j]).startswith("Protein") and str(hdr[j + 1]) == "Label":
        lab = None
        for r in rows[2:]:
            if r[j + 1]: lab = str(r[j + 1]); break
        blocks.append((j, lab))
print("blocks:", blocks, flush=True)

recs = []
for j, lab in blocks:
    for r in rows[2:]:
        if r[j] is None or r[j + 2] is None: continue
        a = str(r[j])
        g = acc2gene.get(a)
        if not g: continue
        try: fc = float(r[j + 2]); p = float(r[j + 3]) if r[j + 3] is not None else np.nan
        except (TypeError, ValueError): continue
        recs.append({"comparison": lab, "gene": g, "log2FC": fc, "p": p})
MS = pd.DataFrame(recs)
MS.to_csv(f"{OUT}/L1e_msstats_long.tsv", sep="\t", index=False)
print(MS.groupby("comparison").size(), flush=True)

# ---------- tests per comparison ----------
asn_q90 = FE.Asn_content.quantile(0.9)
aat_q90 = FE.AAT_freq.quantile(0.9)
out_rows = []
for comp, sub in MS.groupby("comparison"):
    M = sub.merge(FE, left_on="gene", right_index=True, how="inner").dropna(subset=["log2FC"])
    hi_asn = M.Asn_content >= asn_q90
    hi_aat = M.AAT_freq >= aat_q90
    res = {"comparison": comp, "n_prot": len(M),
           "n_AsnTop10_detected": int(hi_asn.sum()),
           "median_log2FC_AsnTop10": float(M.loc[hi_asn, "log2FC"].median()),
           "median_log2FC_rest": float(M.loc[~hi_asn, "log2FC"].median()),
           "MW_p_AsnTop10_up": float(stats.mannwhitneyu(M.loc[hi_asn, "log2FC"], M.loc[~hi_asn, "log2FC"], alternative="greater").pvalue),
           "spearman_AsnContent_log2FC": float(stats.spearmanr(M.Asn_content, M.log2FC).statistic),
           "spearman_p": float(stats.spearmanr(M.Asn_content, M.log2FC).pvalue),
           "MW_p_AATTop10_up": float(stats.mannwhitneyu(M.loc[hi_aat, "log2FC"], M.loc[~hi_aat, "log2FC"], alternative="greater").pvalue),
           "median_log2FC_AP2": float(M.loc[M.is_ap2, "log2FC"].median()) if M.is_ap2.any() else np.nan,
           "n_AP2_detected": int(M.is_ap2.sum()),
           "median_log2FC_PUF": float(M.loc[M.is_puf, "log2FC"].median()) if M.is_puf.any() else np.nan,
           "n_PUF_detected": int(M.is_puf.sum())}
    out_rows.append(res)
    print(comp, ":", {k: round(v, 4) if isinstance(v, float) else v for k, v in res.items()}, flush=True)
R = pd.DataFrame(out_rows)
R.to_csv(f"{OUT}/L1e_protein_tests.tsv", sep="\t", index=False)

qc = {"blocks": [b[1] for b in blocks], "tests": out_rows}
json.dump(qc, open(f"{OUT}/L1e_qc.json", "w"), indent=2, default=str)
open(f"{OUT}/claim_impact.md", "w").write(f"""# L1e claim impact (ring 3: protein layer; exploratory)

- date: 2026-09-04
- data: Small-Saunders 2024 (Nat Microbiol) TMT MSstats source data, Dd2 WT vs R539T, 12h post DHA/DMSO pulse.
- prediction from chain (L1 mRNA up + GSE226632 TE up + Asn supply down): Asn-rich proteins shift under DHA.
- results per comparison in L1e_protein_tests.tsv; chain verdict updated in STATUS after user review.
- caveat: TMT detects ~2-3k proteins (coverage biased to abundant); Asn-rich proteins are often LOW abundance
  (regulatory) -> detection bias against the test set, noted for honest interpretation.
- claim impact: protein-layer evidence for/against the Asn-axis chain; no claim change yet.
""")
mani = pd.DataFrame({"file": [XLSX, CDS, f"{ROOT}/data/raw/veupathdb/PlasmoDB-71/PlasmoDB-71_Pfalciparum3D7.gff"]})
mani["sha256"] = mani.file.map(sha); mani["size"] = mani.file.map(os.path.getsize)
mani.to_csv(f"{OUT}/input_manifest.tsv", sep="\t", index=False)
outs = [f for f in os.listdir(OUT) if f != "checksums.sha256" and os.path.isfile(f"{OUT}/{f}")]
with open(f"{OUT}/checksums.sha256", "w") as fh:
    for f in sorted(outs):
        fh.write(f"{sha(os.path.join(OUT, f))}  {f}\n")
print("DONE", flush=True)
